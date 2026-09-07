import docx
from docx.shared import Pt, Inches
from models import LessonPlan, SchemeOfWork

def create_lesson_plan_docx(lesson_plan: LessonPlan, output_path: str = "Lesson_Plan.docx"):
    """Generates a nicely formatted Word document for a Lesson Plan."""
    doc = docx.Document()
    
    # Title & Metadata
    title = doc.add_heading('CBC Lesson Plan', 0)
    title.alignment = 1 
    
    p = doc.add_paragraph()
    p.add_run("Subject: ").bold = True
    p.add_run(f"{lesson_plan.subject} | ")
    p.add_run("Grade: ").bold = True
    p.add_run(f"{lesson_plan.grade}\n")
    p.add_run("Week: ").bold = True
    p.add_run(f"{lesson_plan.week} | ")
    p.add_run("Lesson: ").bold = True
    p.add_run(f"{lesson_plan.lesson_number}")
    p.alignment = 1
    
    # Strand Table
    table = doc.add_table(rows=2, cols=2)
    table.style = 'Table Grid'
    table.cell(0,0).text = 'Strand'
    table.cell(0,1).text = lesson_plan.strand
    table.cell(1,0).text = 'Sub-Strand'
    table.cell(1,1).text = lesson_plan.sub_strand
    
    doc.add_paragraph()
    
    # Outcomes & Questions
    doc.add_heading('Specific Learning Outcomes', level=2)
    for outcome in lesson_plan.specific_learning_outcomes:
        doc.add_paragraph(outcome, style='List Bullet')
        
    doc.add_heading('Key Inquiry Question', level=2)
    doc.add_paragraph(lesson_plan.key_inquiry_question)
        
    # Organisation of Learning (Intro/Dev/Conclusion)
    doc.add_heading('Organisation of Learning', level=2)
    doc.add_paragraph("Introduction:", style='Body Text').runs[0].bold = True
    doc.add_paragraph(lesson_plan.organisation_of_learning.get('introduction', ''))
    
    doc.add_paragraph("Lesson Development:", style='Body Text').runs[0].bold = True
    doc.add_paragraph(lesson_plan.organisation_of_learning.get('lesson_development', ''))
    
    doc.add_paragraph("Conclusion:", style='Body Text').runs[0].bold = True
    doc.add_paragraph(lesson_plan.organisation_of_learning.get('conclusion', ''))
        
    # Resources & Integration
    doc.add_heading('Learning Resources', level=2)
    for res in lesson_plan.resources:
        doc.add_paragraph(res, style='List Bullet')
        
    doc.add_heading('Integration', level=2)
    doc.add_paragraph(f"Core Competencies: {', '.join(lesson_plan.core_competencies)}")
    doc.add_paragraph(f"Values: {', '.join(lesson_plan.values)}")
    doc.add_paragraph(f"PCIs: {', '.join(lesson_plan.pcis)}")
    
    doc.add_heading('Assessment Methods', level=2)
    for method in lesson_plan.assessment_methods:
        doc.add_paragraph(method, style='List Bullet')

    # Teacher Reflection
    doc.add_heading('Reflection', level=2)
    reflection_text = lesson_plan.reflection if lesson_plan.reflection else "_________________________________________________________________________________\n_________________________________________________________________________________"
    doc.add_paragraph(reflection_text)
    
    doc.save(output_path)
    print(f" Saved Lesson Plan to {output_path}")
    return output_path


def create_scheme_of_work_docx(scheme: SchemeOfWork, output_path: str = "Scheme_of_Work.docx"):
    """Generates a landscape Word document for a Scheme of Work."""
    doc = docx.Document()
    
    # Make document landscape for the wide table
    section = doc.sections[-1]
    new_width, new_height = section.page_height, section.page_width
    section.orientation = docx.enum.section.WD_ORIENT.LANDSCAPE
    section.page_width = new_width
    section.page_height = new_height
    
    title = doc.add_heading(f"Scheme of Work: {scheme.subject} - Grade {scheme.grade}", 0)
    title.alignment = 1
    
    doc.add_paragraph(f"Term: {scheme.term} | Year: {scheme.year}")
    
    # Create the massive Scheme table
    table = doc.add_table(rows=1, cols=7)
    table.style = 'Table Grid'
    
    headers = ['Week', 'Lesson', 'Strand / Sub-Strand', 'Specific Learning Outcomes', 'Learning Experiences', 'Learning Resources', 'Assessment']
    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        for run in hdr_cells[i].paragraphs[0].runs:
            run.font.bold = True
            
    # Populate rows
    for week in scheme.weeks:
        for lesson in week.lessons:
            row_cells = table.add_row().cells
            row_cells[0].text = str(week.week_number)
            row_cells[1].text = str(lesson.lesson_number)
            row_cells[2].text = f"{lesson.strand}\n\n{lesson.sub_strand}"
            
            # Bullet points for outcomes & experiences
            row_cells[3].text = "\n".join([f"• {o}" for o in lesson.specific_learning_outcomes])
            row_cells[4].text = "\n".join([f"• {e}" for e in lesson.learning_experiences])
            row_cells[5].text = "\n".join([f"• {r}" for r in lesson.learning_resources])
            row_cells[6].text = "\n".join([f"• {a}" for a in lesson.assessment_methods])

    doc.save(output_path)
    print(f" Saved Scheme of Work to {output_path}")
    return output_path